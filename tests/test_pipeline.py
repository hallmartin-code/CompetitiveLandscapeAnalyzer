"""End-to-end pipeline behaviour, research gating, and missing-data handling.

These run entirely offline (the session-wide `_offline` fixture strips
credentials), so they exercise the degraded path end to end -- which is exactly
the path most likely to be hit in the field and least likely to be noticed when
it goes wrong.
"""

from __future__ import annotations

import json
from pathlib import Path

import pdfplumber

from models.schemas import Consent, UserInputs
from pipeline import analyse, render, run_analysis
from research.source_manager import SourceManager, evaluate_policy


# --------------------------------------------------------------------------
# Research authorisation -- three independent gates
# --------------------------------------------------------------------------


def test_research_refused_without_consent(settings):
    policy = evaluate_policy(Consent(), settings)
    assert not policy.permitted
    assert "not authorised" in policy.reason


def test_research_refused_when_disabled_by_operator(settings):
    settings.research_enabled = False
    policy = evaluate_policy(Consent().grant(public_research=True), settings)
    assert not policy.permitted
    assert "operator configuration" in policy.reason


def test_research_permitted_only_with_both_gates(settings):
    policy = evaluate_policy(Consent().grant(public_research=True), settings)
    assert policy.permitted
    assert policy.allowed_domains


def test_consent_granted_but_research_declined_is_refused(settings):
    """Authorising the analysis is not the same as authorising web research."""
    consent = Consent().grant(public_research=False)
    assert not evaluate_policy(consent, settings).permitted


def test_domain_allowlist_is_enforced(settings, diagnostics):
    manager = SourceManager(Consent().grant(True), diagnostics, settings)
    assert manager.check_url("https://www.sec.gov/some/filing")
    assert not manager.check_url("https://random-blog.example.com/post")
    assert any("not on the approved source list" in w for w in diagnostics.warnings)


def test_refused_urls_are_recorded(settings, diagnostics):
    manager = SourceManager(Consent().grant(True), diagnostics, settings)
    manager.check_url("https://not-allowed.example.com")
    assert manager.refused


def test_urls_are_not_fetched_without_consent(settings, diagnostics):
    """The gate must hold before any network call is attempted."""
    manager = SourceManager(Consent(), diagnostics, settings)
    assert manager.fetch_urls(["https://www.sec.gov/filing"]) == []
    assert manager.refused


def test_policy_reason_is_always_recorded(settings, diagnostics):
    SourceManager(Consent(), diagnostics, settings)
    assert diagnostics.info, "the research decision must be visible in diagnostics"


def test_transmission_gate_is_independent_of_research(settings, diagnostics):
    manager = SourceManager(Consent().grant(True), diagnostics, settings)
    may, reason = manager.may_transmit()
    assert not may and "not authorised" in reason


# --------------------------------------------------------------------------
# End-to-end, offline
# --------------------------------------------------------------------------


def test_full_run_without_a_model_still_delivers_both_artefacts(
    sample_decks, settings, tmp_path
):
    analysis, pdf_path, json_path = run_analysis(
        sample_decks["pptx"],
        consent=Consent().grant(public_research=False),
        settings=settings,
        output_dir=tmp_path,
    )
    assert pdf_path and pdf_path.exists()
    assert json_path and json_path.exists()
    with pdfplumber.open(str(pdf_path)) as pdf:
        assert len(pdf.pages) == 1


def test_degraded_run_is_flagged_everywhere_it_matters(sample_decks, settings, tmp_path):
    """A reader must not be able to mistake a degraded run for a full one."""
    analysis, pdf_path, _ = run_analysis(
        sample_decks["pptx"], settings=settings, output_dir=tmp_path
    )
    assert analysis.diagnostics.degraded
    assert any("without a language model" in w for w in analysis.diagnostics.warnings)

    with pdfplumber.open(str(pdf_path)) as pdf:
        text = pdf.pages[0].extract_text() or ""
    assert "DEGRADED RUN" in text.upper()


def test_run_records_that_no_research_happened(sample_decks, settings, tmp_path):
    analysis, _, json_path = run_analysis(
        sample_decks["pptx"], settings=settings, output_dir=tmp_path
    )
    record = json.loads(Path(json_path).read_text(encoding="utf-8"))
    assert record["run"]["research_performed"] is False
    assert record["consent"]["public_research_permitted"] is False


def test_user_inputs_flow_through_to_the_record(sample_decks, settings, tmp_path):
    user = UserInputs(
        user_name="Hall Martin", user_email="analyst@example.com",
        company_name="Renamed Health", sector="Digital Health",
    )
    _, _, json_path = run_analysis(
        sample_decks["pptx"], user=user, settings=settings, output_dir=tmp_path
    )
    record = json.loads(Path(json_path).read_text(encoding="utf-8"))
    assert record["user"]["name"] == "Hall Martin"
    assert record["company"]["name"] == "Renamed Health"
    assert record["company"]["sector"] == "Digital Health"


def test_pdf_and_json_share_a_stem(sample_decks, settings, tmp_path):
    _, pdf_path, json_path = run_analysis(
        sample_decks["pptx"], settings=settings, output_dir=tmp_path
    )
    assert pdf_path.stem == json_path.stem
    assert "kestrel" in pdf_path.stem


# --------------------------------------------------------------------------
# Failure handling
# --------------------------------------------------------------------------


def test_missing_deck_returns_errors_rather_than_raising(settings, tmp_path):
    analysis, pdf_path, json_path = run_analysis(
        tmp_path / "nope.pdf", settings=settings, output_dir=tmp_path
    )
    assert pdf_path is None and json_path is None
    assert analysis.diagnostics.errors


def test_corrupt_deck_returns_an_actionable_error(settings, tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"not a pdf at all")
    analysis, pdf_path, _ = run_analysis(bad, settings=settings, output_dir=tmp_path)
    assert pdf_path is None
    assert any("corrupted" in e or "Could not read" in e for e in analysis.diagnostics.errors)


def test_unsupported_deck_type_is_rejected_clearly(settings, tmp_path):
    odd = tmp_path / "deck.key"
    odd.write_bytes(b"keynote")
    analysis, pdf_path, _ = run_analysis(odd, settings=settings, output_dir=tmp_path)
    assert pdf_path is None
    assert any("PDF or PPTX" in e for e in analysis.diagnostics.errors)


def test_image_only_deck_produces_a_report_that_says_so(settings, tmp_path):
    """We still deliver, but the reader is told the evidence base is empty."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas as pdfcanvas

    path = tmp_path / "images.pdf"
    c = pdfcanvas.Canvas(str(path), pagesize=letter)
    for _ in range(5):
        c.rect(80, 80, 300, 300, fill=1)
        c.showPage()
    c.save()

    analysis = analyse(path, settings=settings)
    assert any("no extractable text" in e for e in analysis.diagnostics.errors)
    assert analysis.profile.coverage == 0.0

    pdf_path, _ = render(analysis, tmp_path, settings=settings)
    with pdfplumber.open(str(pdf_path)) as pdf:
        assert len(pdf.pages) == 1          # still exactly one page


def test_analysis_never_raises_on_a_valid_but_bare_deck(settings, tmp_path):
    """A deck with almost no content must degrade, not explode."""
    from pptx import Presentation

    path = tmp_path / "bare.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])      # blank layout
    prs.save(str(path))

    analysis = analyse(path, settings=settings)
    assert analysis.diagnostics.errors or analysis.diagnostics.warnings


def test_supporting_documents_are_ingested(sample_decks, settings, tmp_path):
    extra = tmp_path / "market_note.txt"
    extra.write_text(
        "CMS-0057-F requires payer prior authorization APIs by January 2027.",
        encoding="utf-8",
    )
    analysis = analyse(
        sample_decks["pptx"], supporting_paths=[extra], settings=settings
    )
    assert analysis.profile.company_name.known


def test_progress_callback_reports_every_step(sample_decks, settings):
    seen: list = []
    analyse(
        sample_decks["pptx"], settings=settings,
        progress=lambda i, n, label: seen.append((i, n, label)),
    )
    assert len(seen) == 10
    assert [i for i, _, _ in seen] == list(range(1, 11))
    assert all(n == 10 for _, n, _ in seen)
